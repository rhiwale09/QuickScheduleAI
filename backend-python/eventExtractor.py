"""
QuickScheduleAI — Python Event Extraction Service
==================================================
Handles all document types and text formats.

Endpoints:
  GET  /health
  POST /extract-text   { "text": "..." }     -> { "events": [...] }
  POST /extract-file   multipart file upload -> { "events": [...] }
"""

import re
import io
import os
import logging
import subprocess
import tempfile
from datetime import datetime, timedelta
from typing import Optional
from flask import Flask, request, jsonify
import dateparser

app = Flask(__name__)
logging.basicConfig(level=logging.INFO)
log = logging.getLogger(__name__)

# ── CONSTANTS ────────────────────────────────────────────────────

MONTH_MAP = {
    "january":1,"february":2,"march":3,"april":4,"may":5,"june":6,
    "july":7,"august":8,"september":9,"october":10,"november":11,"december":12,
    "jan":1,"feb":2,"mar":3,"apr":4,"jun":6,"jul":7,"aug":8,
    "sep":9,"sept":9,"oct":10,"nov":11,"dec":12,
}

SEASON_MAP = {
    "early fall":9,"fall":10,"late fall":11,
    "early spring":3,"spring":4,"late spring":5,
    "early summer":6,"summer":7,"late summer":8,"winter":1,
}

EVENT_KEYWORDS = re.compile(
    r"\b(meeting|dinner|party|prom|graduation|chapel|ceremony|reception|"
    r"baccalaureate|homecoming|assembly|dance|skit|brunch|banquet|conference|"
    r"presentation|orientation|webinar|fair|drive|election|exam|test|deadline|"
    r"due|celebration|sunrise|sunset|tea|luncheon|breakfast|lunch|session|"
    r"workshop|retreat|concert|performance|game|tournament|commencement|"
    r"convocation|showcase|expo|gala|parade|rally|competition|holiday|"
    r"dismissal|break|begins|program|seminar|registration|forms|faculty|"
    r"school|class|student)\b",
    re.IGNORECASE
)

SKIP_RE = re.compile(
    r"\b(submit|upload|download|purchase|format|criteria|handbook|jostens|"
    r"jpeg|sign up will|please include|link to|naviance|registration deadline|"
    r"late registration|please visit|click here|learn more)\b",
    re.IGNORECASE
)

MONTH_HEADER_RE = re.compile(
    r"^(JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST|SEPTEMBER|"
    r"OCTOBER|NOVEMBER|DECEMBER)\s*(202\d)?$",
    re.IGNORECASE
)

DAY_PREFIX_RE = re.compile(
    r"^(monday|tuesday|wednesday|thursday|friday|saturday|sunday)[,.]?\s*",
    re.IGNORECASE
)

FULL_DATE_RE = re.compile(
    r"\b(january|february|march|april|may|june|july|august|september|"
    r"october|november|december)\s+\d{1,2}(?:st|nd|rd|th)?(?:,?\s*202\d)?\b",
    re.IGNORECASE
)

NUMERIC_DATE_RE = re.compile(r"\b\d{1,2}/\d{1,2}\b")

TIME_RANGE_RE = re.compile(
    r"(\d{1,2})(?::(\d{2}))?\s*(a\.?m\.?|p\.?m\.?)?\s*[-\u2013]\s*"
    r"(\d{1,2})(?::(\d{2}))?\s*(a\.?m\.?|p\.?m\.?)",
    re.IGNORECASE
)

SINGLE_TIME_RE = re.compile(
    r"\b(\d{1,2})(?::(\d{2}))?\s*(a\.?m\.?|p\.?m\.?)\b",
    re.IGNORECASE
)

# ── DATE HELPERS ─────────────────────────────────────────────────

def detect_ref_year(text):
    years = [int(m) for m in re.findall(r"\b(202[4-9])\b", text)]
    return min(years) if years else datetime.now().year + 1


def apply_time(text, base):
    m = TIME_RANGE_RE.search(text)
    if m:
        sh = int(m.group(1))
        sm_min = int(m.group(2) or 0)
        eh = int(m.group(4))
        em_min = int(m.group(5) or 0)
        sap = (m.group(3) or "").lower().replace(".", "")
        eap = (m.group(6) or "").lower().replace(".", "")
        if eap == "pm" and eh < 12:
            eh += 12
        if sap == "pm" and sh < 12:
            sh += 12
        if eap == "pm" and not sap and sh < eh < 12:
            sh += 12
        if sap == "am" and sh == 12:
            sh = 0
        start = base.replace(hour=sh, minute=sm_min, second=0, microsecond=0)
        end = base.replace(hour=eh, minute=em_min, second=0, microsecond=0)
        return start, end, True

    m2 = SINGLE_TIME_RE.search(text)
    if m2:
        h = int(m2.group(1))
        mn = int(m2.group(2) or 0)
        ap = m2.group(3).lower().replace(".", "")
        if ap == "pm" and h < 12:
            h += 12
        if ap == "am" and h == 12:
            h = 0
        start = base.replace(hour=h, minute=mn, second=0, microsecond=0)
        end = start + timedelta(hours=1)
        return start, end, True

    return base, base + timedelta(hours=1), False


def parse_full_date(text, ref_year):
    text = DAY_PREFIX_RE.sub("", text).strip()
    m = re.search(
        r"(january|february|march|april|may|june|july|august|september|"
        r"october|november|december)\s+(\d{1,2})(?:st|nd|rd|th)?(?:,?\s*(202\d))?",
        text, re.IGNORECASE
    )
    if not m:
        return None
    month = MONTH_MAP[m.group(1).lower()]
    day = int(m.group(2))
    year = int(m.group(3)) if m.group(3) else ref_year
    try:
        return datetime(year, month, day)
    except ValueError:
        return None


def parse_numeric_date(m_str, d_str, y_str, ref_year, ctx_month=None):
    month, day = int(m_str), int(d_str)
    if not (1 <= month <= 12 and 1 <= day <= 31):
        return None
    if y_str:
        year = 2000 + int(y_str) if len(y_str) == 2 else int(y_str)
    else:
        year = ref_year
        if ctx_month and month < ctx_month:
            year = ref_year + 1
    try:
        return datetime(year, month, day)
    except ValueError:
        return None


def dateparser_fallback(text, ref_year):
    try:
        settings = {
            "PREFER_DAY_OF_MONTH": "first",
            "PREFER_DATES_FROM": "future",
            "RETURN_AS_TIMEZONE_AWARE": False,
        }
        result = dateparser.parse(text, settings=settings)
        if result and abs(result.year - ref_year) <= 2:
            return result
    except Exception:
        pass
    return None


def season_date(text, ref_year):
    lower = text.lower()
    for season, month in SEASON_MAP.items():
        if season in lower:
            return datetime(ref_year, month, 1)
    for name, month in MONTH_MAP.items():
        if len(name) > 3 and re.search(rf"\b{name}\b", lower):
            return datetime(ref_year, month, 1)
    return None


# ── TITLE CLEANER ────────────────────────────────────────────────

def clean_title(text):
    if not text:
        return ""
    months = ("JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST|"
              "SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER")
    text = re.sub(rf"^(?:{months})(?:\s+202\d)?[\t ]+", "", text, flags=re.IGNORECASE)
    text = re.sub(
        r"\s+ON\s+(?:MONDAY|TUESDAY|WEDNESDAY|THURSDAY|FRIDAY|SATURDAY|SUNDAY)?,?\s*(?:"
        + months + r")\s+\d+", "", text, flags=re.IGNORECASE
    )
    text = re.sub(
        r"\s+DUE\s+(?:MONDAY|TUESDAY|WEDNESDAY|THURSDAY|FRIDAY|SATURDAY|SUNDAY)?,?\s*(?:"
        + months + r")\s+\d+", " (due)", text, flags=re.IGNORECASE
    )
    text = re.sub(r"\s*[-\u2013:]\s*$", "", text)
    text = re.sub(r"^[-\u2013:\s]+", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:80]


def to_title_case(text):
    skip = {"a","an","the","and","or","but","in","on","at","to","for","of","with","by"}
    words = text.lower().split()
    return " ".join(
        w.capitalize() if i == 0 or w not in skip else w
        for i, w in enumerate(words)
    )


def make_event(title, start, end, has_time, description="", confidence=80):
    return {
        "title":       clean_title(title)[:80],
        "start":       start.isoformat(),
        "end":         end.isoformat(),
        "location":    "",
        "description": description[:200],
        "ambiguous":   not has_time,
        "recurrence":  None,
        "confidence":  confidence,
    }


# ── STRATEGY A: Tab-separated table parser ───────────────────────

def parse_table(text, ref_year):
    events = []
    for line in text.split("\n"):
        line = line.strip()
        if "\t" not in line or len(line) < 5:
            continue
        parts = [p.strip() for p in line.split("\t") if p.strip()]
        if len(parts) < 2:
            continue

        # Date in first column
        base = parse_full_date(parts[0], ref_year)
        if base and len(parts[1]) > 2:
            title = clean_title(parts[1])
            if len(title) < 3 or SKIP_RE.search(title):
                continue
            start, end, has_time = apply_time(parts[1], base)
            events.append(make_event(title, start, end, has_time, line, 95))
            continue

        # Title in col 0, date in col 1 (test tables)
        if len(parts) >= 2:
            base2 = parse_full_date(parts[1], ref_year)
            if base2 and 1 < len(parts[0]) < 20:
                if re.match(r"^(test|date|program|deadline|registration)",
                            parts[0], re.IGNORECASE):
                    continue
                start, end, has_time = apply_time(parts[1], base2)
                events.append(make_event(parts[0], start, end, has_time, line, 90))

    return events


# ── STRATEGY B: Calendar format parser ──────────────────────────

def parse_calendar(text, ref_year):
    events = []
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    ctx_month = None
    ctx_year = ref_year
    pending = None

    def year_for(month):
        if not ctx_month or month >= ctx_month:
            return ctx_year
        return ctx_year + 1

    i = 0
    while i < len(lines):
        line = lines[i]

        hm = MONTH_HEADER_RE.match(line)
        if hm:
            ctx_month = MONTH_MAP[hm.group(1).lower()]
            if hm.group(2):
                ctx_year = int(hm.group(2))
            pending = None
            i += 1
            continue

        if SKIP_RE.search(line) or len(line) < 3:
            i += 1
            continue

        stripped = DAY_PREFIX_RE.sub("", line).strip()

        # Cross-month range: "12/21 - 1/3/27  Winter Break"
        rm = re.match(
            r"(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?\s*[-\u2013]\s*"
            r"(?:(?:monday|tuesday|wednesday|thursday|friday|saturday|sunday),?\s*)?"
            r"(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?\s{2,}(.*)",
            line, re.IGNORECASE
        )
        if rm:
            sm, sd = int(rm.group(1)), int(rm.group(2))
            em, ed = int(rm.group(4)), int(rm.group(5))
            sy = (2000 + int(rm.group(3)) if rm.group(3) and len(rm.group(3)) == 2
                  else int(rm.group(3)) if rm.group(3) else year_for(sm))
            ey = (2000 + int(rm.group(6)) if rm.group(6) and len(rm.group(6)) == 2
                  else int(rm.group(6)) if rm.group(6)
                  else (ctx_year + 1 if em < sm else ctx_year))
            title = clean_title(rm.group(7).strip())
            if title and len(title) > 1:
                try:
                    start = datetime(sy, sm, sd)
                    end = datetime(ey, em, ed)
                    events.append(make_event(title, start, end, False, title, 90))
                    pending = None
                except ValueError:
                    pass
            i += 1
            continue

        # Date + title 2+ spaces: "1/4/27   Second Semester Begins"
        nt = re.match(r"^(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?\s{2,}(.+)", stripped)
        if nt:
            base = parse_numeric_date(
                nt.group(1), nt.group(2), nt.group(3),
                year_for(int(nt.group(1))), ctx_month
            )
            if base:
                title = clean_title(nt.group(4).strip())
                if title and len(title) > 1:
                    s, e, ht = apply_time(nt.group(4), base)
                    events.append(make_event(title, s, e, ht, title, 90))
                    pending = None
                i += 1
                continue

        # Pure numeric date
        pn = re.match(r"^(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?\s*$", stripped)
        if pn:
            base = parse_numeric_date(
                pn.group(1), pn.group(2), pn.group(3),
                year_for(int(pn.group(1))), ctx_month
            )
            if base:
                next_line = lines[i + 1] if i + 1 < len(lines) else ""
                next_ok = (next_line and not MONTH_HEADER_RE.match(next_line)
                           and not NUMERIC_DATE_RE.search(next_line)
                           and 2 < len(next_line) < 100)
                if pending:
                    title = clean_title(pending)
                    events.append(make_event(
                        title, base, base + timedelta(hours=1), False, title, 90))
                    pending = None
                elif next_ok:
                    title = clean_title(next_line.strip())
                    events.append(make_event(
                        title, base, base + timedelta(hours=1), False, title, 90))
                    i += 1
            i += 1
            continue

        # Short gap date+title: "9/7  School Holiday"
        sn = re.match(r"^(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?\s+(.+)", stripped)
        if sn and len(sn.group(4).strip()) > 2:
            base = parse_numeric_date(
                sn.group(1), sn.group(2), sn.group(3),
                year_for(int(sn.group(1))), ctx_month
            )
            if base:
                title = clean_title(sn.group(4).strip())
                events.append(make_event(
                    title, base, base + timedelta(hours=1), False, title, 90))
                pending = None
                i += 1
                continue

        # Pure title
        has_date = NUMERIC_DATE_RE.search(line) or FULL_DATE_RE.search(line)
        if not has_date and 3 < len(line) < 120:
            pending = line.strip()

        i += 1

    return events


# ── STRATEGY C: Inline date parser ──────────────────────────────

def parse_inline(text, ref_year):
    events = []
    seen = set()
    lines = [re.sub(r"^[●○■•\-\*\s\t]+", "", l).strip()
             for l in text.split("\n") if l.strip()]

    mp = (r"(?:january|february|march|april|may|june|july|august|"
          r"september|october|november|december)")
    dp = r"(?:monday|tuesday|wednesday|thursday|friday|saturday|sunday)"

    SESSION_RE = re.compile(
        rf"^(Session\s+\w+)\s*:\s*({mp}\s+\d{{1,2}}(?:\s*[-\u2013]\s*\d{{1,2}})?)"
        rf"\s*\(([^)]+)\)", re.IGNORECASE
    )
    HEADING_RE = re.compile(
        rf"^([A-Z][A-Z\s&\/,]+?)\s+(?:ON|DUE(?:\s+{dp})?(?:,)?)\s+"
        rf"(?:{dp},?\s+)?({mp}\s+\d{{1,2}})", re.IGNORECASE
    )
    INLINE_A = re.compile(
        rf"^(?:{dp},?\s+)?({mp}\s+\d{{1,2}}(?:st|nd|rd|th)?"
        rf"(?:,?\s*202\d)?)\s*[-\u2013:(),]+\s*(.+)", re.IGNORECASE
    )
    INLINE_B = re.compile(
        rf"^(.+?)\s+(?:on\s+)?({mp}\s+\d{{1,2}}(?:st|nd|rd|th)?"
        rf"(?:,?\s*202\d)?)[.,]?\s*$", re.IGNORECASE
    )
    STANDALONE = re.compile(
        rf"^(?:{dp},?\s+)?({mp}\s+\d{{1,2}}(?:st|nd|rd|th)?"
        rf"(?:,?\s*202\d)?)\s*(?:\(([^)]+)\))?\s*([-\u2013:]\s*.{{0,60}})?$",
        re.IGNORECASE
    )

    for line in lines:
        if SKIP_RE.search(line) or len(line) < 8 or "\t" in line:
            continue

        title = date_str = time_str = None

        sm = SESSION_RE.match(line)
        if sm:
            title = sm.group(1).strip()
            date_str = sm.group(2).strip()
            time_str = sm.group(3).strip()

        if not title:
            hm = HEADING_RE.match(line)
            if hm:
                title = clean_title(hm.group(1).strip())
                date_str = hm.group(2).strip()

        if not title:
            sm2 = STANDALONE.match(line)
            if sm2:
                date_str = sm2.group(1)
                time_str = sm2.group(2) or ""
                rest = (sm2.group(3) or "").strip().lstrip("-\u2013: ")
                title = rest if rest and len(rest) > 2 else None
                if not title:
                    date_str = None

        if not title:
            ma = INLINE_A.match(line)
            if ma:
                date_str = ma.group(1)
                title = ma.group(2).strip()

        if not title:
            mb = INLINE_B.match(line)
            if mb:
                title = mb.group(1).strip()
                date_str = mb.group(2)

        if not date_str or not title or len(title) < 3:
            continue

        title = clean_title(title)
        if len(title) < 3 or SKIP_RE.search(title):
            continue

        base = dateparser_fallback(f"{date_str} {ref_year}", ref_year)
        if not base:
            base = parse_full_date(date_str, ref_year)
        if not base:
            continue

        time_text = time_str if time_str else line
        start, end, has_time = apply_time(time_text, base)

        key = f"{title.lower()[:20]}|{start.date()}"
        if key in seen:
            continue
        seen.add(key)

        events.append(make_event(title, start, end, has_time, line[:200], 80))

    return events


# ── STRATEGY D: Block/paragraph parser ──────────────────────────

def classify_block(text):
    score = 0.0
    if FULL_DATE_RE.search(text):        score += 0.35
    if re.search(r"\b202[4-9]\b", text): score += 0.10
    if SINGLE_TIME_RE.search(text):      score += 0.15
    if re.match(r"[A-Z][a-z]+ \d{1,2}[\s]*[-\u2013:]", text): score += 0.20
    first = text.split("\n")[0].strip()
    if (first == first.upper() and 4 < len(first) < 70
            and re.search(r"[A-Z]{2,}", first)):
        score += 0.20
    if EVENT_KEYWORDS.search(text): score += 0.08
    if SKIP_RE.search(text):        score -= 0.45
    return max(0.0, min(1.0, score))


def extract_block_title(text):
    lines = [re.sub(r"^[●○■•\-\*\s]+", "", l).strip()
             for l in text.split("\n") if l.strip()]
    for line in lines:
        c = re.sub(r"\s*[-\u2013\u2014]\s*(student only event|parent only|continued.*)",
                   "", line, flags=re.IGNORECASE).strip()
        if c == c.upper() and 4 < len(c) < 70 and re.search(r"[A-Z]{2,}", c):
            return clean_title(to_title_case(c))

    im = re.search(
        r"(?:january|february|march|april|may|june|july|august|september|"
        r"october|november|december)\s+\d{1,2}(?:st|nd|rd|th)?(?:\s*[-\u2013]\s*\d+)?"
        r"(?:\s*,?\s*202\d)?(?:\s*\([^)]*\))?\s*[-\u2013:(]+\s*(.+)",
        text, re.IGNORECASE
    )
    if im:
        t = clean_title(im.group(1))
        if len(t) > 3:
            return t

    for line in lines:
        if 5 < len(line) < 80:
            return clean_title(line)

    return clean_title(text[:60])


def parse_blocks(text, ref_year):
    events = []
    blocks = [b.strip() for b in re.split(r"\n{2,}", text)
              if 20 < len(b.strip()) < 1500]

    for block in blocks:
        if SKIP_RE.search(block):
            continue
        score = classify_block(block)
        if score < 0.25:
            continue

        base = None
        date_match = FULL_DATE_RE.search(block)
        if date_match:
            base = dateparser_fallback(
                date_match.group(0) + f" {ref_year}", ref_year)
        if not base:
            base = parse_full_date(block, ref_year)
        if not base:
            base = season_date(block, ref_year)
        if not base:
            continue

        title = extract_block_title(block)
        if not title or len(title) < 3:
            continue

        start, end, has_time = apply_time(block, base)
        clean_block = " ".join(block.split())
        first_sent = re.match(r"^[^.!?]+[.!?]", clean_block)
        desc = (first_sent.group(0) if first_sent else clean_block)[:200]

        events.append(make_event(title, start, end, has_time, desc, int(score * 100)))

    return events


# ── DEDUPLICATOR ─────────────────────────────────────────────────

def dice_sim(a, b):
    if a == b: return 1.0
    if len(a) < 2 or len(b) < 2: return 0.0
    bg = {}
    for i in range(len(a) - 1):
        s = a[i:i+2]
        bg[s] = bg.get(s, 0) + 1
    x = 0
    for i in range(len(b) - 1):
        s = b[i:i+2]
        if bg.get(s, 0) > 0:
            x += 1
            bg[s] -= 1
    return (2 * x) / (len(a) + len(b) - 2)


def deduplicate(events):
    seen = {}
    result = []
    for e in events:
        dk = datetime.fromisoformat(e["start"]).date().isoformat()
        tk = re.sub(r"[^a-z0-9]", "", e["title"].lower())[:20]
        key = f"{tk}|{dk}"
        if key in seen:
            continue
        duplicate = False
        for k in seen:
            if k.endswith(dk):
                other_tk = k.split("|")[0]
                if dice_sim(tk, other_tk) > 0.75:
                    duplicate = True
                    break
        if duplicate:
            continue
        seen[key] = True
        result.append(e)
    return result


# ── FILE TEXT EXTRACTORS ─────────────────────────────────────────

def extract_text_from_pdf(file_bytes):
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages = [p.extract_text() or "" for p in pdf.pages]
            text = "\n\n".join(pages)
            if len(text.strip()) > 50:
                return text
    except Exception as e:
        log.warning(f"pdfplumber failed: {e}")
    try:
        from pdf2image import convert_from_bytes
        import pytesseract
        images = convert_from_bytes(file_bytes)
        return "\n\n".join(pytesseract.image_to_string(img) for img in images)
    except Exception as e:
        log.warning(f"PDF OCR failed: {e}")
    return ""


def extract_text_from_docx(file_bytes):
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        log.warning(f"DOCX extraction failed: {e}")
    return ""


def extract_text_from_doc(file_bytes):
    try:
        with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name
        result = subprocess.run(
            ["antiword", tmp_path],
            capture_output=True, text=True, timeout=10
        )
        os.unlink(tmp_path)
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout
    except Exception as e:
        log.warning(f"antiword failed: {e}")
    try:
        text = file_bytes.decode("latin-1", errors="ignore")
        lines = []
        for line in text.split("\n"):
            clean = "".join(c for c in line if 32 <= ord(c) < 127 or c == "\t")
            if len(clean.strip()) > 5:
                lines.append(clean.strip())
        return "\n".join(lines)
    except Exception:
        return ""


def extract_text_from_xlsx(file_bytes):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(
            io.BytesIO(file_bytes), read_only=True, data_only=True)
        rows = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row
                         if c is not None and str(c).strip()]
                if cells:
                    rows.append("\t".join(cells))
        return "\n".join(rows)
    except Exception as e:
        log.warning(f"openpyxl failed: {e}")
    try:
        import xlrd
        wb = xlrd.open_workbook(file_contents=file_bytes)
        rows = []
        for sheet in wb.sheets():
            for rx in range(sheet.nrows):
                cells = [str(sheet.cell_value(rx, cx))
                         for cx in range(sheet.ncols)
                         if str(sheet.cell_value(rx, cx)).strip()]
                if cells:
                    rows.append("\t".join(cells))
        return "\n".join(rows)
    except Exception as e:
        log.warning(f"xlrd failed: {e}")
    return ""


def extract_text_from_csv(file_bytes):
    import csv
    for enc in ["utf-8", "latin-1", "cp1252"]:
        try:
            text = file_bytes.decode(enc)
            reader = csv.reader(io.StringIO(text))
            return "\n".join("\t".join(row) for row in reader if any(row))
        except Exception:
            continue
    return ""


def extract_text_from_eml(file_bytes):
    import email
    from email import policy as email_policy
    try:
        msg = email.message_from_bytes(
            file_bytes, policy=email_policy.default)
        parts = []
        if msg["subject"]:
            parts.append(f"Subject: {msg['subject']}")
        if msg["date"]:
            parts.append(f"Date: {msg['date']}")
        if msg.is_multipart():
            for part in msg.walk():
                ct = part.get_content_type()
                if ct == "text/plain":
                    parts.append(
                        part.get_payload(decode=True).decode(
                            "utf-8", errors="ignore"))
                elif ct == "text/html":
                    html = part.get_payload(decode=True).decode(
                        "utf-8", errors="ignore")
                    parts.append(re.sub(r"<[^>]+>", " ", html))
        else:
            payload = msg.get_payload(decode=True)
            if payload:
                parts.append(payload.decode("utf-8", errors="ignore"))
        return "\n\n".join(parts)
    except Exception as e:
        log.warning(f"EML extraction failed: {e}")
        return extract_text_from_txt(file_bytes)


def extract_text_from_image(file_bytes):
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(io.BytesIO(file_bytes))
        return pytesseract.image_to_string(img)
    except Exception as e:
        log.warning(f"Image OCR failed: {e}")
    return ""


def extract_text_from_txt(file_bytes):
    for enc in ["utf-8", "latin-1", "cp1252"]:
        try:
            text = file_bytes.decode(enc)
            text = re.sub(r"<[^>]+>", " ", text)        # strip HTML
            text = re.sub(r"\\[a-z]+\d*\s?", " ", text) # strip RTF
            text = re.sub(r"[{}]", " ", text)
            return text
        except Exception:
            continue
    return ""


# ── MAIN EXTRACTOR ───────────────────────────────────────────────

def extract_events(text, max_events=60):
    ref_year = detect_ref_year(text)
    log.info(f"ref_year={ref_year} text_len={len(text)}")

    has_tabs = "\t" in text
    has_cal  = bool(MONTH_HEADER_RE.search(text))
    has_num  = bool(NUMERIC_DATE_RE.search(text))

    events = []
    if has_tabs:
        a = parse_table(text, ref_year)
        log.info(f"Strategy A (table): {len(a)}")
        events.extend(a)
    if has_cal or has_num:
        b = parse_calendar(text, ref_year)
        log.info(f"Strategy B (calendar): {len(b)}")
        events.extend(b)
    c = parse_inline(text, ref_year)
    log.info(f"Strategy C (inline): {len(c)}")
    events.extend(c)
    d = parse_blocks(text, ref_year)
    log.info(f"Strategy D (blocks): {len(d)}")
    events.extend(d)

    unique = deduplicate(events)
    unique.sort(key=lambda e: (e["start"], -e["confidence"]))
    for i, e in enumerate(unique[:max_events]):
        e["id"] = i + 1

    log.info(f"Final: {len(unique[:max_events])} events")
    return unique[:max_events]


# ── FLASK ENDPOINTS ──────────────────────────────────────────────

@app.route("/health", methods=["GET"])
def health():
    return jsonify({
        "status": "ok",
        "service": "QuickScheduleAI Python Extractor"
    })


@app.route("/extract-text", methods=["POST"])
def extract_text_endpoint():
    data = request.get_json()
    if not data or "text" not in data:
        return jsonify({"error": "Missing 'text' field"}), 400
    try:
        events = extract_events(data["text"])
        return jsonify({"events": events, "count": len(events)})
    except Exception as e:
        log.error(f"Extraction error: {e}")
        return jsonify({"error": str(e)}), 500


@app.route("/extract-file", methods=["POST"])
def extract_file_endpoint():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    f = request.files["file"]
    name = f.filename.lower()
    data = f.read()

    try:
        if name.endswith(".pdf"):
            text = extract_text_from_pdf(data)
        elif name.endswith(".docx"):
            text = extract_text_from_docx(data)
        elif name.endswith(".doc"):
            text = extract_text_from_doc(data)
            if not text.strip():
                text = extract_text_from_docx(data)
        elif name.endswith((".xlsx", ".xls")):
            text = extract_text_from_xlsx(data)
        elif name.endswith(".csv"):
            text = extract_text_from_csv(data)
        elif name.endswith(".eml"):
            text = extract_text_from_eml(data)
        elif name.endswith((".html", ".htm", ".rtf")):
            text = extract_text_from_txt(data)
        elif any(name.endswith(x) for x in
                 [".png",".jpg",".jpeg",".tiff",".bmp",".gif",".webp"]):
            text = extract_text_from_image(data)
        else:
            text = extract_text_from_txt(data)
            if len(text.strip()) < 20:
                text = extract_text_from_image(data)

        if not text.strip():
            return jsonify({"error": "Could not extract text from file"}), 400

        events = extract_events(text)
        return jsonify({
            "events": events,
            "count": len(events),
            "raw_text": text[:2000]
        })

    except Exception as e:
        log.error(f"File extraction error: {e}")
        return jsonify({"error": str(e)}), 500


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    app.run(host="0.0.0.0", port=port, debug=False)
